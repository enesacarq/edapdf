
from pathlib import Path
import fitz


def pdf_to_images(input_path, output_dir, image_format="png"):
    # Kontrollojmë formatin e imazhit
    image_format = image_format.lower()

    if image_format not in ["png", "jpg"]:
        raise ValueError("Format must be PNG or JPG")

    # Krijojmë folderin e rezultateve
    output_folder = Path(output_dir)
    output_folder.mkdir(parents=True, exist_ok=True)

    # Hapim dokumentin PDF
    document = fitz.open(input_path)
    image_paths = []

    try:
        # Përpunojmë çdo faqe të PDF-së
        for page_number, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(dpi=200, alpha=False)

            image_path = output_folder / f"page_{page_number}.{image_format}"

            pixmap.save(str(image_path))
            image_paths.append(str(image_path))

    finally:
        document.close()

    return image_paths


def images_to_pdf(input_paths, output_path):
    from PIL import Image

    if not input_paths:
        raise ValueError("No images provided")

    images = []

    try:
        for image_path in input_paths:
            with Image.open(image_path) as img:
                images.append(img.convert("RGB").copy())

        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)

        images[0].save(
            str(output_file),
            "PDF",
            save_all=True,
            append_images=images[1:]
        )

        return str(output_file)

    finally:
        for image in images:
            image.close()

            
def merge_pdfs(input_paths, output_path):
    if not input_paths:
        raise ValueError("No PDF files provided")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    merged_document = fitz.open()

    try:
        for pdf_path in input_paths:
            with fitz.open(pdf_path) as document:
                merged_document.insert_pdf(document)

        merged_document.save(str(output_file))

    finally:
        merged_document.close()

    return str(output_file)



def split_pdf(input_path, page_ranges, output_dir):
    if not page_ranges:
        raise ValueError("No page ranges provided")

    output_folder = Path(output_dir)
    output_folder.mkdir(parents=True, exist_ok=True)

    created_files = []

    with fitz.open(input_path) as document:
        total_pages = len(document)

        for index, (start, end) in enumerate(page_ranges, start=1):
            if start < 1 or end > total_pages or start > end:
                raise ValueError("Invalid page range")

            new_document = fitz.open()

            try:
                new_document.insert_pdf(
                    document,
                    from_page=start - 1,
                    to_page=end - 1
                )

                output_path = output_folder / f"part_{index}.pdf"
                new_document.save(str(output_path))
                created_files.append(str(output_path))

            finally:
                new_document.close()

    return created_files



def delete_pdf_pages(input_path, pages, output_path):
    if not pages:
        raise ValueError("No pages provided")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with fitz.open(input_path) as document:
        total_pages = len(document)

        for page_number in pages:
            if page_number < 1 or page_number > total_pages:
                raise ValueError("Invalid page number")

        if len(set(pages)) >= total_pages:
            raise ValueError("Cannot delete all pages")

        # Fshijmë faqet duke filluar nga numri më i madh
        for page_number in sorted(set(pages), reverse=True):
            document.delete_page(page_number - 1)

        document.save(str(output_file))

    return str(output_file)



def add_watermark(input_path, output_path, text):
    if not text or not text.strip():
        raise ValueError("Watermark text cannot be empty")

    if Path(input_path).resolve() == Path(output_path).resolve():
        raise ValueError("Input and output paths must be different")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with fitz.open(input_path) as document:
        for page in document:
            page_rect = page.rect

            font_size = 40
            text_width = fitz.get_text_length(
                text,
                fontname="helv",
                fontsize=font_size
            )

            x = (page_rect.width - text_width) / 2
            y = page_rect.height / 2

            page.insert_text(
                (x, y),
                text,
                fontsize=font_size,
                fontname="helv",
                color=(0.6, 0.6, 0.6),
                fill_opacity=0.3,
                overlay=True
            )

        document.save(str(output_file))

    return str(output_file)



def pdf_to_word(input_path, output_path):
    from pdf2docx import Converter

    if not Path(input_path).is_file():
        raise FileNotFoundError("PDF file not found")

    if Path(input_path).resolve() == Path(output_path).resolve():
        raise ValueError("Input and output paths must be different")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    converter = Converter(str(input_path))

    try:
        converter.convert(str(output_file))
    finally:
        converter.close()

    return str(output_file)



def pdf_to_powerpoint(input_path, output_path):
    from pptx import Presentation
    from pptx.util import Inches
    from io import BytesIO

    if not Path(input_path).is_file():
        raise FileNotFoundError("PDF file not found")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()

    with fitz.open(input_path) as document:
        if len(document) == 0:
            raise ValueError("PDF has no pages")

        first_page = document[0]
        first_rect = first_page.rect

        presentation.slide_width = Inches(first_rect.width / 72)
        presentation.slide_height = Inches(first_rect.height / 72)

        blank_layout = presentation.slide_layouts[6]

        for page in document:
            slide = presentation.slides.add_slide(blank_layout)

            pixmap = page.get_pixmap(
                dpi=200,
                alpha=False
            )

            image_stream = BytesIO(pixmap.tobytes("png"))

            page_width = Inches(page.rect.width / 72)
            page_height = Inches(page.rect.height / 72)

            scale = min(
                presentation.slide_width / page_width,
                presentation.slide_height / page_height
            )

            image_width = int(page_width * scale)
            image_height = int(page_height * scale)

            left = (presentation.slide_width - image_width) // 2
            top = (presentation.slide_height - image_height) // 2

            slide.shapes.add_picture(
                image_stream,
                left,
                top,
                width=image_width,
                height=image_height
            )

    presentation.save(str(output_file))

    return str(output_file)



def pdf_to_excel(input_path, output_path):
    import pdfplumber
    from openpyxl import Workbook

    if not Path(input_path).is_file():
        raise FileNotFoundError("PDF file not found")

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PDF Data"

    current_row = 1
    tables_found = 0

    with pdfplumber.open(input_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()

            for table in tables:
                if not table:
                    continue

                tables_found += 1

                worksheet.cell(
                    row=current_row,
                    column=1,
                    value=f"Page {page_number} - Table {tables_found}"
                )
                current_row += 1

                for row in table:
                    for column_number, cell_value in enumerate(row, start=1):
                        worksheet.cell(
                            row=current_row,
                            column=column_number,
                            value=cell_value
                        )

                    current_row += 1

                current_row += 2

    if tables_found == 0:
        worksheet.cell(
            row=1,
            column=1,
            value="No tables found in PDF"
        )

    workbook.save(str(output_file))

    return str(output_file)



def office_to_pdf(input_path, output_path):
    import subprocess
    import tempfile
    import shutil

    input_file = Path(input_path).resolve()
    output_file = Path(output_path).resolve()

    soffice_path = Path(
        r"C:\Program Files\LibreOffice\program\soffice.exe"
    )

    if not input_file.is_file():
        raise FileNotFoundError("Office file not found")

    if input_file.suffix.lower() not in [".docx", ".xlsx", ".pptx"]:
        raise ValueError("Unsupported Office format")

    if not soffice_path.is_file():
        raise FileNotFoundError("LibreOffice not found")

    output_file.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as temp_dir:
        result = subprocess.run(
            [
                str(soffice_path),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(input_file)
            ],
            capture_output=True,
            text=True,
            timeout=120
        )

        converted_pdf = Path(temp_dir) / (
            input_file.stem + ".pdf"
        )

        if result.returncode != 0 or not converted_pdf.exists():
            raise RuntimeError(
                "Office conversion failed: "
                + result.stderr
                + result.stdout
            )

        shutil.move(str(converted_pdf), str(output_file))

    return str(output_file)
